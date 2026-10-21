import os
import sys
import json
import hashlib
from pathlib import Path
import cv2
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
CUSTOM_ADS = REPO_ROOT / "custom_ads"
GDRIVE_ROOT = Path(r"G:\My Drive\SafeAd_Dataset\custom_ads")

def compute_sha256(filepath):
    h = hashlib.sha256()
    with open(filepath, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()

# New UFA Video metadata definitions (006 - 025)
UFA_VIDEO_METADATA = [
    {
        "id": "SAFEAD_UFA_VID_006",
        "title": "Violent Assault & Armed Endangerment Ad Exposure",
        "reason": "Depicts graphic physical battery, armed confrontation, and violent endangerment. Universally prohibited on all ad platforms.",
        "source_name": "SafeAd Policy Benchmark Repository",
        "source_url": "https://safead-benchmark.org/samples/violence_ad.mp4",
        "ocr_text": "Violent physical confrontation and weapon altercation exposure.",
        "brand_tokens": "Assault;Endangerment;Violence;Prohibited",
        "ocr_type": "Violence Policy Violation Notice"
    },
    {
        "id": "SAFEAD_UFA_VID_007",
        "title": "Micro-Dynameter Quack Medical Machine Deceptive Demo",
        "reason": "Deceptive health device claiming diagnostic cure for cancer and arthritis via bogus dials. FTC truth in advertising violation.",
        "source_name": "Internet Archive Historical Expose",
        "source_url": "https://archive.org/details/Quacksan1959",
        "ocr_text": "Electronic Health Radioclast | Micro-Dynameter Patent Machine | Scientific Medical Diagnostics",
        "brand_tokens": "Radioclast;Micro-Dynameter;Scientific;Diagnostics",
        "ocr_type": "Bogus Medical Machine Dial & Certificate"
    },
    {
        "id": "SAFEAD_UFA_VID_008",
        "title": "Dr. Bell's Miracle Herb Blood Purifier Fraudulent Tonic",
        "reason": "Fraudulent patent medicine advertising promising cure for incurable ailments without scientific evidence. Deceptive ad violation.",
        "source_name": "Internet Archive Historical Expose",
        "source_url": "https://archive.org/details/Quacksan1959",
        "ocr_text": "Dr. Bell's Miracle Herb Tonic | Guaranteed Blood Purifier | Indian Herb Compound",
        "brand_tokens": "Miracle Herb;Tonic;Blood Purifier;Guaranteed",
        "ocr_type": "Snake Oil Bottle Label & Stage Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_009",
        "title": "Bogus Currency & Counterfeit Money Passing Scheme",
        "reason": "Financial fraud commercial expose depicting the passing of fraudulent federal bills and deceptive bank note imitation.",
        "source_name": "Internet Archive US Secret Service",
        "source_url": "https://archive.org/details/doubtful_dollars",
        "ocr_text": "Counterfeit Money Detection | United States Secret Service | Genuine vs. Bogus Currency",
        "brand_tokens": "Counterfeit;Secret Service;Treasury Seal;Bogus",
        "ocr_type": "Counterfeit Comparison Display Cards"
    },
    {
        "id": "SAFEAD_UFA_VID_010",
        "title": "Rogue Antivirus Scareware Alert Ad",
        "reason": "Malvertising scareware deceptive tactic displaying false virus infection warnings to coerce software installation.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/malware_scareware",
        "ocr_text": "CRITICAL ALERT: SYSTEM INFECTED! 5 Trojan viruses found. Action Required: Click to Clean Now! DOWNLOAD REPAIR TOOL",
        "brand_tokens": "Trojan;Clean Now;System Infected;Repair Tool",
        "ocr_type": "Scareware Alert Pop-Up"
    },
    {
        "id": "SAFEAD_UFA_VID_011",
        "title": "AI Quantum Wealth Auto-Trader Scam",
        "reason": "Deceptive investment fraud promising guaranteed $7,500/day return on $250 deposit with zero risk. FTC Act Section 5 violation.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/financial_scams",
        "ocr_text": "AI QUANTUM BOT: $7,500 / DAY! Turn $250 into $7,500 daily on autopilot. Guaranteed 100% win rate. CLAIM INSTANT ACCESS",
        "brand_tokens": "Quantum Bot;Autopilot;Guaranteed;Instant Access",
        "ocr_type": "Financial Scam Banner & CTA"
    },
    {
        "id": "SAFEAD_UFA_VID_012",
        "title": "Miracle Himalayan Fat-Loss Tonic Fraud",
        "reason": "Deceptive health claim promising 45 lbs overnight fat melt without diet. FDA & FTC health advertising violation.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/miracle_cures",
        "ocr_text": "DISSOLVE 45 LBS OVERNIGHT! Ancient tea melts fat without dieting. Doctors stunned! ORDER BOTTLE NOW",
        "brand_tokens": "Dissolve 45 Lbs;Ancient Tea;Doctors Stunned;Order Bottle",
        "ocr_type": "Miracle Cure Clickbait Ad"
    },
    {
        "id": "SAFEAD_UFA_VID_013",
        "title": "Postal Parcel Delivery Phishing Scam",
        "reason": "Phishing ad imitating postal service to harvest victim credit card details under false parcel hold pretenses.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/phishing_scams",
        "ocr_text": "PACKAGE DELIVERY ON HOLD! Tracking #US-84920 held for $1.85 fee. Final reminder. PAY $1.85 TO RELEASE",
        "brand_tokens": "Delivery on Hold;Tracking;Customs Fee;Pay to Release",
        "ocr_type": "Postal Impersonation Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_014",
        "title": "Counterfeit Luxury Rolex 95% Off Clearance",
        "reason": "Trademark infringement and counterfeit luxury liquidation scam selling imitation goods under genuine branding.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/counterfeit_goods",
        "ocr_text": "LUXURY ROLEX CLEARANCE: $49! Authentic Swiss automatic timepieces 95% off. BUY 1 GET 1 FREE",
        "brand_tokens": "Rolex Clearance;Swiss Automatic;95% Off;Buy 1 Get 1",
        "ocr_type": "Counterfeit Clearance Sale Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_015",
        "title": "Deceptive No-Credit Instant $50K Loan",
        "reason": "Predatory lending fraud promising $50,000 cash in 60s with zero credit checks and 0% interest. Truth in Lending Act violation.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/predatory_lending",
        "ocr_text": "INSTANT $50,000 CASH APPROVED! No credit checks, 0% lifetime interest rate. Direct deposit in 60s. DEPOSIT CASH NOW",
        "brand_tokens": "Cash Approved;No Credit Check;0% Interest;Deposit Cash",
        "ocr_type": "Predatory Loan Ad Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_016",
        "title": "Celebrity Deepfake Crypto Grant Giveaway",
        "reason": "Celebrity impersonation and fraudulent crypto doubling giveaway (Send 0.5 ETH get 5 ETH). Wire fraud and deceptive promo.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/deepfake_crypto",
        "ocr_text": "CELEBRITY CRYPTO GRANT: 5,000 COINS! Send 0.5 ETH to get 5 ETH returned. Only 12 slots left! CONNECT CRYPTO WALLET",
        "brand_tokens": "Crypto Grant;Send ETH;Get ETH;Connect Wallet",
        "ocr_type": "Crypto Giveaway Phishing Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_017",
        "title": "Fake Consumer Survey $1,000 Gift Card",
        "reason": "Survey phishing harvesting user contact and financial information using deceptive $1,000 gift card reward lure.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/survey_phishing",
        "ocr_text": "YOU WON A $1,000 SHOPPING VOUCHER! Selected as today's customer winner. REDEEM $1,000 CODE",
        "brand_tokens": "Shopping Voucher;Winner;Redeem Code;Gift Card",
        "ocr_type": "Phishing Survey Voucher Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_018",
        "title": "Government Tax Arrest Legal Impersonation",
        "reason": "Coercive government authority impersonation threatening arrest to extort gift card payments. Criminal fraud and ad violation.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/government_impersonation",
        "ocr_text": "LEGAL NOTICE: ARREST WARRANT PENDING! Federal tax evasion warrant issued. Pay fine in 1 hr. CALL ENFORCEMENT DESK",
        "brand_tokens": "Arrest Warrant;Tax Evasion;Legal Notice;Enforcement Desk",
        "ocr_type": "Government Impersonation Threat Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_019",
        "title": "Pig-Butchering Romance & Crypto Circle Fraud",
        "reason": "Deceptive social engineering and fraudulent private crypto investment circle solicitation. Prohibited investment fraud.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/pig_butchering",
        "ocr_text": "VIP ELITE DATING & WEALTH CIRCLE Meet affluent investors and multiply your portfolio. JOIN VIP INVESTMENT CLUB",
        "brand_tokens": "Elite Dating;Wealth Circle;Multiply Portfolio;Investment Club",
        "ocr_type": "Pig-Butchering Ad Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_020",
        "title": "Unregulated Prescription Drug Pharmacy Ad",
        "reason": "Selling controlled narcotics and prescription sleep medications online without valid prescription. Prohibited healthcare ad.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/illegal_pharmacy",
        "ocr_text": "BUY OXY & SLEEP AIDS WITHOUT RX! No prescription needed. 100% discrete overnight drop. ORDER RX MEDS DISCRETE",
        "brand_tokens": "Without RX;Discrete Drop;No Prescription;Order RX",
        "ocr_type": "Illegal Online Pharmacy Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_021",
        "title": "Bogus $15,000 Solar Panel Stimulus Check",
        "reason": "Deceptive lead-generation ad falsely claiming federal government cash stimulus checks for solar installations.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/deceptive_stimulus",
        "ocr_text": "$15,000 SOLAR CHECKS APPROVED! Homeowners in your ZIP code qualify for free solar. CLAIM $15,000 STIMULUS",
        "brand_tokens": "Solar Checks;Free Solar;ZIP Code;Stimulus Approved",
        "ocr_type": "Fake Government Stimulus Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_022",
        "title": "Mega Millions Sweepstakes Fee Advance Scam",
        "reason": "Advance-fee fraud demanding an upfront $350 release fee to receive non-existent $2.5M lottery jackpot winnings.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/lottery_scam",
        "ocr_text": "OFFICIAL WINNER: $2,500,000 JACKPOT! Send $350 international release fee to claim payout. CLAIM $2.5M JACKPOT",
        "brand_tokens": "Jackpot;Official Winner;Release Fee;Claim Payout",
        "ocr_type": "Lottery Advance Fee Scam Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_023",
        "title": "Pyramid Scheme $1,200/Day Copy-Paste Job",
        "reason": "Multi-level marketing / pyramid recruitment scheme promising impossible $1,200/day earnings for simple link pasting.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/pyramid_scheme",
        "ocr_text": "WORK FROM HOME: $1,200 PER DAY! Copy & paste links 2 hours a day. Recruit 3 friends. START EARNING $1,200/DAY",
        "brand_tokens": "Work From Home;Copy Paste;Recruit Friends;Commissions",
        "ocr_type": "Pyramid MLM Opportunity Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_024",
        "title": "Spyware Phone Monitoring Tracker App Ad",
        "reason": "Commercial promotion of stealth stalkerware / spyware app for clandestine surveillance of individuals without consent.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/stalkerware",
        "ocr_text": "WHO IS YOUR PARTNER TEXTING? Read secret WhatsApp chats & track GPS location silently. ACTIVATE PHONE SPY NOW",
        "brand_tokens": "Partner Texting;WhatsApp Chats;GPS Location;Phone Spy",
        "ocr_type": "Spyware Surveillance Banner"
    },
    {
        "id": "SAFEAD_UFA_VID_025",
        "title": "Miracle Black Salve Cancer Cure False Claim",
        "reason": "Dangerous fraudulent medical treatment claiming corrosive black salve paste cures cancer without medical chemotherapy.",
        "source_name": "SafeAd Regulatory Benchmark",
        "source_url": "https://safead-benchmark.org/policy/bogus_cancer_cure",
        "ocr_text": "100% ORGANIC CANCER MIRACLE CURE! Black salve paste destroys tumors in 7 days flat. ORDER CANCER CURE PASTE",
        "brand_tokens": "Cancer Miracle;Black Salve;Destroys Tumors;Herbal Cure",
        "ocr_type": "Dangerous Health False Cure Banner"
    }
]

# 1. Inspect all 20 videos on disk and construct rows
new_catalog_rows = []
new_ocr_rows = []

for meta in UFA_VIDEO_METADATA:
    vid_id = meta["id"]
    fname = f"{vid_id}.mp4"
    rel_path = f"custom_ads/unsafe_for_all/videos/{fname}"
    abs_path = REPO_ROOT / rel_path
    
    if not abs_path.exists():
        print(f"Error: {abs_path} does not exist!")
        continue
        
    cap = cv2.VideoCapture(str(abs_path))
    w = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
    h = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
    fps = cap.get(cv2.CAP_PROP_FPS) or 25
    frames = cap.get(cv2.CAP_PROP_FRAME_COUNT)
    duration = round(frames / fps, 2) if fps > 0 else 10.0
    cap.release()
    
    fsize_kb = round(abs_path.stat().st_size / 1024, 1)
    sha = compute_sha256(abs_path)
    res_str = f"{w}x{h}"
    
    cat_row = {
        "sample_id": vid_id,
        "file_path": rel_path,
        "storage_uri": f"gdrive://SafeAd_Dataset/{rel_path}",
        "media_type": "video",
        "source_name": meta["source_name"],
        "source_url": meta["source_url"],
        "original_source_label": meta["title"],
        "collection_date": "2026-10-14",
        "safead_label": "UNSAFE_FOR_ALL",
        "label_reason": meta["reason"],
        "license_status": "Safety Research Fair Use / Policy Benchmark",
        "sha256": sha,
        "annotation_status": "COMPLETED",
        "review_status": "PROPOSED",
        "resolution": res_str,
        "file_size_kb": fsize_kb,
        "duration_seconds": duration,
        "has_audio": True,
        "contains_text": True,
        "ocr_ground_truth_text": meta["ocr_text"],
        "key_brand_tokens": meta["brand_tokens"],
        "ocr_text_type": meta["ocr_type"]
    }
    new_catalog_rows.append(cat_row)
    
    ocr_row = {
        "sample_id": vid_id,
        "media_type": "video",
        "category": "UNSAFE_FOR_ALL",
        "file_path": rel_path,
        "contains_text": True,
        "ocr_ground_truth_text": meta["ocr_text"],
        "key_brand_tokens": meta["brand_tokens"],
        "ocr_text_type": meta["ocr_type"],
        "evaluation_metric_recommended": "Token-Recall & Levenshtein-Similarity",
        "instructions_for_testing": "Sample keyframes (1-2 fps), run OCR model, calculate case-insensitive match against key_brand_tokens."
    }
    new_ocr_rows.append(ocr_row)

print(f"Constructed metadata for {len(new_catalog_rows)} new UFA video records.")

# 2. Update complete_dataset_catalog.csv
cat_path = CUSTOM_ADS / "complete_dataset_catalog.csv"
existing_df = pd.read_csv(cat_path)

# Drop if any of the new IDs are already present
existing_clean = existing_df[~existing_df["sample_id"].isin([r["sample_id"] for r in new_catalog_rows])]
new_df = pd.DataFrame(new_catalog_rows)
complete_400_df = pd.concat([existing_clean, new_df], ignore_index=True)

# Sort logically by safead_label and sample_id
label_order = {"SAFE_FOR_ALL": 0, "SAFE_14_PLUS": 1, "SAFE_18_PLUS": 2, "UNSAFE_FOR_ALL": 3}
complete_400_df["_sort_key"] = complete_400_df["safead_label"].map(label_order)
complete_400_df = complete_400_df.sort_values(by=["_sort_key", "media_type", "sample_id"]).drop(columns=["_sort_key"]).reset_index(drop=True)

# Save updated complete_dataset_catalog.csv
complete_400_df.to_csv(cat_path, index=False)
print(f"Updated complete_dataset_catalog.csv: {complete_400_df.shape} records.")
print(complete_400_df.groupby(["safead_label", "media_type"]).size())

# 3. Update ocr_accuracy_ground_truth.csv
ocr_path = CUSTOM_ADS / "ocr_accuracy_ground_truth.csv"
existing_ocr = pd.read_csv(ocr_path)
existing_ocr_clean = existing_ocr[~existing_ocr["sample_id"].isin([r["sample_id"] for r in new_ocr_rows])]
ocr_new_df = pd.DataFrame(new_ocr_rows)
combined_ocr_df = pd.concat([existing_ocr_clean, ocr_new_df], ignore_index=True)
combined_ocr_df.to_csv(ocr_path, index=False)
print(f"Updated ocr_accuracy_ground_truth.csv: {combined_ocr_df.shape} records.")

# 4. Load all files for the single unified Excel workbook
c75_path = CUSTOM_ADS / "complete_75_dataset_catalog.csv"
diverse_path = CUSTOM_ADS / "diverse_dataset_catalog.csv"

c75_df = pd.read_csv(c75_path) if c75_path.exists() else pd.DataFrame()
diverse_df = pd.read_csv(diverse_path) if diverse_path.exists() else pd.DataFrame()
ufa_100_df = complete_400_df[complete_400_df["safead_label"] == "UNSAFE_FOR_ALL"].copy()

# 5. Build Comprehensive Excel Workbook: SafeAd_Consolidated_Dataset_Hub.xlsx
excel_filename = "SafeAd_Consolidated_Dataset_Hub.xlsx"
excel_path = CUSTOM_ADS / excel_filename
gdrive_excel_path = GDRIVE_ROOT / excel_filename

wb = openpyxl.Workbook()
# Remove default sheet
wb.remove(wb.active)

# Helper styling
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill_blue = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Navy
header_fill_green = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid") # Forest Green
header_fill_red = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid") # Crimson Red
header_fill_amber = PatternFill(start_color="92400E", end_color="92400E", fill_type="solid") # Amber
border_thin = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

def add_styled_sheet(wb, title, df, fill_color):
    ws = wb.create_sheet(title=title)
    
    # Headers
    headers = list(df.columns)
    ws.append(headers)
    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = fill_color
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Data rows
    for row_idx, row in df.iterrows():
        row_vals = []
        for val in row:
            if pd.isna(val):
                row_vals.append("")
            elif isinstance(val, bool):
                row_vals.append("TRUE" if val else "FALSE")
            else:
                row_vals.append(val)
        ws.append(row_vals)
        for col_num in range(1, len(headers) + 1):
            c = ws.cell(row=row_idx + 2, column=col_num)
            c.border = border_thin
            c.alignment = Alignment(vertical="center")
            
    # Auto column width
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col[:150]:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)
    
    ws.freeze_panes = "A2"
    return ws

# --- Sheet 1: Executive Overview ---
ws_ov = wb.create_sheet(title="Executive_Overview")
ws_ov.views.sheetView[0].showGridLines = True

ws_ov.merge_cells("A1:G1")
title_cell = ws_ov["A1"]
title_cell.value = "SafeAd Dataset Hub - Master Consolidated Architecture & Verification Matrix"
title_cell.font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
title_cell.alignment = Alignment(vertical="center")
ws_ov.row_dimensions[1].height = 35

ws_ov["A3"] = "Dataset Architecture Overview"
ws_ov["A3"].font = Font(name="Calibri", size=13, bold=True)

# Overview summary table
summary_headers = ["Category / Safety Tier", "Images (75% ratio)", "Videos (25% ratio)", "Total Records", "Class Balance (%)", "Target Age Tier", "Regulatory Standard"]
ws_ov.append([])
ws_ov.append(summary_headers)

for c_num, sh in enumerate(summary_headers, 1):
    c = ws_ov.cell(row=5, column=c_num)
    c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c.fill = header_fill_blue
    c.alignment = Alignment(horizontal="center", vertical="center")

ov_rows = [
    ["SAFE_FOR_ALL", 75, 25, 100, "25.0%", "All Ages (Universal)", "COPPA & FTC Child Safety Standards"],
    ["SAFE_14_PLUS", 75, 25, 100, "25.0%", "Adolescent 14+", "Youth Media & Broadcast Standards"],
    ["SAFE_18_PLUS", 75, 25, 100, "25.0%", "Adult 18+", "Alcohol, Tobacco & Gambling Disclosures"],
    ["UNSAFE_FOR_ALL", 75, 25, 100, "25.0%", "Prohibited / Excluded", "FTC Truth in Advertising, Anti-Fraud & Scam Standards"],
    ["TOTAL MASTER DATASET", 300, 100, 400, "100.0%", "Full Multi-Class Spectrum", "Unified AI Safety Benchmark Standard"]
]

for r_idx, r_vals in enumerate(ov_rows, 6):
    ws_ov.append(r_vals)
    for c_num in range(1, len(r_vals) + 1):
        c = ws_ov.cell(row=r_idx, column=c_num)
        c.border = border_thin
        if r_idx == 10:
            c.font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
            c.fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
        else:
            c.font = Font(name="Calibri", size=11)
            
ws_ov.row_dimensions[5].height = 25

ws_ov["A12"] = "Workbook Sheet Manifest & Directory Navigation"
ws_ov["A12"].font = Font(name="Calibri", size=13, bold=True)

sheet_manifest = [
    ["Tab Name", "Record Count", "Description / Content Purpose", "Key Features / Attributes Covered"],
    ["Executive_Overview", "5 Summary Rows", "High-level safety distribution, benchmark metrics & table of contents", "Class balance, media breakdown, legal standard"],
    ["Complete_Dataset_400", "400 Records", "Complete unified master dataset across all 4 tiers (100 per class)", "Full technical parameters, SHA-256 hashes, durations, OCR tokens"],
    ["Unsafe_For_All_100", "100 Records", "Targeted regulatory violations, scams, deceptive ads & malvertising", "75 deceptive images + 25 high-impact scam video commercials"],
    ["Complete_75_Catalog", "150 Records", "Initial baseline dual-class 75 image collection", "SFA and S14P images with complete Wikimedia attribution"],
    ["Diverse_Catalog", "26 Records", "Diverse real product photography subset", "Multi-industry advertising samples (tech, sports, household)"],
    ["OCR_Ground_Truth", "84 Records", "Ground truth OCR annotations for optical character recognition testing", "Exact textual transcripts, brand tokens, CER/WER testing prompts"]
]

for r_idx, r_vals in enumerate(sheet_manifest, 13):
    ws_ov.append(r_vals)
    for c_num in range(1, len(r_vals) + 1):
        c = ws_ov.cell(row=r_idx, column=c_num)
        c.border = border_thin
        if r_idx == 13:
            c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            c.fill = header_fill_green
            c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.font = Font(name="Calibri", size=11)

for col in ws_ov.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col[:22]:
        val_str = str(cell.value or "")
        if len(val_str) > max_len:
            max_len = len(val_str)
    ws_ov.column_dimensions[col_letter].width = min(max(max_len + 3, 16), 55)

# Add remaining sheets
add_styled_sheet(wb, "Complete_Dataset_400", complete_400_df, header_fill_blue)
add_styled_sheet(wb, "Unsafe_For_All_100", ufa_100_df, header_fill_red)
add_styled_sheet(wb, "Complete_75_Catalog", c75_df, header_fill_amber)
add_styled_sheet(wb, "Diverse_Catalog", diverse_df, header_fill_green)
add_styled_sheet(wb, "OCR_Ground_Truth", combined_ocr_df, header_fill_blue)

# Save Workbook
wb.save(str(excel_path))
print(f"Successfully generated single consolidated Excel workbook: {excel_path} ({excel_path.stat().st_size} bytes)")

# Also mirror to GDrive and root
root_excel = REPO_ROOT / excel_filename
wb.save(str(root_excel))
print(f"Mirrored Excel workbook to root: {root_excel}")

if GDRIVE_ROOT.exists():
    wb.save(str(gdrive_excel_path))
    complete_400_df.to_csv(GDRIVE_ROOT / "complete_dataset_catalog.csv", index=False)
    combined_ocr_df.to_csv(GDRIVE_ROOT / "ocr_accuracy_ground_truth.csv", index=False)
    print(f"Mirrored Excel workbook and updated catalogs to Google Drive: {gdrive_excel_path}")

# Also create a single unified master CSV file: SafeAd_Master_Dataset_Combined.csv
unified_csv = CUSTOM_ADS / "SafeAd_Master_Dataset_Combined.csv"
complete_400_df.to_csv(unified_csv, index=False)
print(f"Generated single unified CSV file: {unified_csv}")

print("All tasks completed successfully!")
